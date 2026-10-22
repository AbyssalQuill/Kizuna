#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文档解析与 Excel 生成辅助工具：
1. PDF 提取（pdftotext + pypdf 双引擎）
2. Excel 提取（openpyxl -> Markdown 表格）
3. PPTX 提取（zip -> slide xml -> 文本）
4. Excel 生成（根据 JSON 生成排版优美的 .xlsx）
"""

import sys
import os
import json
import zipfile
import subprocess
import xml.etree.ElementTree as ET

def parse_pdf(file_path):
    text_parts = []
    # 优先使用 pdftotext
    try:
        res = subprocess.run(["pdftotext", "-layout", file_path, "-"], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=15)
        if res.returncode == 0 and res.stdout:
            txt = res.stdout.decode("utf-8", errors="replace").strip()
            if txt:
                return txt
    except Exception:
        pass

    # 兜底使用 pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        total_pages = len(reader.pages)
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(f"--- 第 {i+1}/{total_pages} 页 ---\n{page_text.strip()}")
        if text_parts:
            return "\n\n".join(text_parts)
    except Exception as e:
        return f"（PDF 解析错误: {e}）"

    return "（PDF 文档内容为空或未能识别出文本）"

def parse_excel(file_path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        sheets_out = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            rows = []
            max_r = 50
            count = 0
            for row in ws.iter_rows(values_only=True):
                count += 1
                if count <= max_r:
                    # 过滤纯空行
                    cleaned = ["" if v is None else str(v).replace("\n", " ").strip() for v in row]
                    if any(cleaned):
                        rows.append(cleaned)
            if not rows:
                continue

            # 对齐列数
            max_cols = max(len(r) for r in rows)
            padded = [r + [""] * (max_cols - len(r)) for r in rows]

            headers = padded[0]
            header_line = "| " + " | ".join(h if h else f"列{j+1}" for j, h in enumerate(headers)) + " |"
            sep_line = "| " + " | ".join(["---"] * max_cols) + " |"
            data_lines = ["| " + " | ".join(r) + " |" for r in padded[1:]]

            sheet_md = f"### 工作表: {sheetname} (共 {count} 行, {max_cols} 列)\n" + header_line + "\n" + sep_line
            if data_lines:
                sheet_md += "\n" + "\n".join(data_lines)
            if count > max_r:
                sheet_md += f"\n\n…（本表共 {count} 行，仅展示前 {max_r} 行）"
            sheets_out.append(sheet_md)

        wb.close()
        if sheets_out:
            return "\n\n".join(sheets_out)
        return "（Excel 表格内容为空）"
    except Exception as e:
        return f"（Excel 解析错误: {e}）"

def parse_pptx(file_path):
    try:
        with zipfile.ZipFile(file_path, 'r') as z:
            slides = [f for f in z.namelist() if f.startswith('ppt/slides/slide') and f.endswith('.xml')]
            # 按幻灯片序号排序
            slides.sort(key=lambda x: int(''.join(filter(str.isdigit, x)) or 0))
            pages = []
            for idx, slide_path in enumerate(slides, 1):
                xml_data = z.read(slide_path)
                root = ET.fromstring(xml_data)
                texts = [elem.text for elem in root.iter() if elem.tag.endswith('}t') and elem.text]
                clean_texts = [t.strip() for t in texts if t.strip()]
                if clean_texts:
                    pages.append(f"--- 幻灯片 {idx} ---\n" + "\n".join(clean_texts))
            if pages:
                return "\n\n".join(pages)
            return "（PPT 幻灯片未提取到文本）"
    except Exception as e:
        return f"（PPT 解析错误: {e}）"

def generate_excel(spec_json, output_path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    # 移除默认 sheet
    default_sheet = wb.active

    sheets = spec_json.get("sheets")
    if not sheets and ("headers" in spec_json or "rows" in spec_json):
        sheets = [{
            "name": spec_json.get("title") or "Sheet1",
            "headers": spec_json.get("headers", []),
            "rows": spec_json.get("rows", [])
        }]

    if not sheets:
        sheets = [{"name": "Sheet1", "headers": ["序号", "项目"], "rows": [["1", "示例数据"]]}]

    # 样式定义
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid") # 优雅蓝
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Microsoft YaHei", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    for s_idx, s in enumerate(sheets):
        ws = wb.create_sheet(title=str(s.get("name") or f"Sheet{s_idx+1}")[:31])
        headers = s.get("headers") or []
        rows = s.get("rows") or []

        # 写入表头
        if headers:
            ws.append(headers)
            for col_idx, cell in enumerate(ws[1], 1):
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
            ws.row_dimensions[1].height = 28

        # 写入数据行
        for r_idx, r in enumerate(rows, start=2 if headers else 1):
            ws.append(r)
            row_cells = ws[r_idx]
            for col_idx, cell in enumerate(row_cells, 1):
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if r_idx % 2 == 1:
                    cell.fill = zebra_fill
            ws.row_dimensions[r_idx].height = 22

        # 自动列宽
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                # 中文字符长度折算
                char_len = sum(2 if ord(c) > 127 else 1 for c in val)
                if char_len > max_len:
                    max_len = char_len
            ws.column_dimensions[col_letter].width = max(10, min(max_len + 3, 50))

    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    wb.save(output_path)
    return {"ok": True, "path": output_path}

def main():
    if len(sys.argv) < 3:
        print(json.dumps({"ok": False, "error": "参数不足"}))
        sys.exit(1)

    cmd = sys.argv[1]
    target = sys.argv[2]

    if cmd == "parse":
        ext = os.path.splitext(target)[1].lower()
        if ext == ".pdf":
            txt = parse_pdf(target)
        elif ext in [".xlsx", ".xls"]:
            txt = parse_excel(target)
        elif ext in [".pptx"]:
            txt = parse_pptx(target)
        else:
            txt = f"不支持的格式: {ext}"
        print(json.dumps({"ok": True, "text": txt}))

    elif cmd == "generate_excel":
        out_file = sys.argv[3] if len(sys.argv) > 3 else "output.xlsx"
        try:
            with open(target, 'r', encoding='utf-8') as f:
                spec = json.load(f)
            res = generate_excel(spec, out_file)
            print(json.dumps(res))
        except Exception as e:
            print(json.dumps({"ok": False, "error": str(e)}))

if __name__ == "__main__":
    main()
